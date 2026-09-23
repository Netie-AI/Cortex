"""EPIC-GRANT-04 (#205): smooth reads on what the operator already granted.

Three reaches, all read-only, all fail-closed:

1. **Window attach** (``attach_window``). A window grant (EPIC-GRANT-01, kind
   ``window``: title + pid as UACC ``list_windows`` names it) is *verified*
   against a live ``list_windows`` result and a read handle (the metadata UACC
   returned for that window) is handed back. The only UACC tool this path may
   call is ``list_windows`` (``ATTACH_TOOL``). Nothing here launches, starts,
   opens or activates an application: if the window is not open right now the
   attach is refused with the reason and the operator opens it themselves
   (R-0015). The master switch and arming are read through ``policy.decide``
   exactly as any capture tool; nothing here arms a server or flips the switch.

   Title match rule: both titles are ``casefold``-ed, surrounding whitespace is
   stripped and internal whitespace runs collapse to one space; they must then
   be **equal**. A pid must match exactly. A document that renamed its window
   since the grant (``Book1`` became ``Book1.xlsx``) is refused so the operator
   grants the window they now see, instead of Crew guessing which window they
   meant. On a plain-text ``list_windows`` result (no JSON) the line must carry
   the pid as a whole number token and contain the normalised title.

2. **Excel as data** (``ws_read_xlsx``). A granted ``.xlsx`` / ``.xlsm`` is
   opened with ``openpyxl`` in ``read_only`` + ``data_only`` mode and rendered as
   tab-separated rows. No COM, no Excel window, no UI driving. Size-capped at
   ``MAX_XLSX_BYTES`` the way ``workspace.MAX_BYTES`` caps a text read.

3. **Browser history** is refused in ``workspace.py`` (``is_history_database``
   here is the one rule both share): there is no grant kind for it in v1.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any

RULE_WINDOW = "R-0015"
RULE_JAIL = "R-0011"

#: The only UACC tool the attach path may call. Pinned by a test.
ATTACH_TOOL = "list_windows"
#: UACC server name as catalogued in ``mcp_client.DEFAULT_SPECS``.
ATTACH_SERVER = "uacc"

#: Cap for a workbook read as data. Workbooks are zipped XML, so this admits a
#: few tens of thousands of cells while keeping a stray export from being
#: streamed into a transcript.
MAX_XLSX_BYTES = 2 * 1024 * 1024
XLSX_ROWS = 200
XLSX_MAX_ROWS = 2000
XLSX_MAX_COLS = 64
XLSX_EXTENSIONS = frozenset({".xlsx", ".xlsm"})

# Chrome / Edge keep visited URLs in a sqlite file literally named ``History``
# (with sqlite's -journal / -wal / -shm sidecars and the older
# ``Archived History``); Firefox keeps them in ``places.sqlite``. No grant kind
# covers these in v1, so a folder grant that happens to contain one must not
# admit it.
_HISTORY_DB_RE = re.compile(
    r"^(?:history|archived history|places\.sqlite)(?:-journal|-wal|-shm)?$",
    re.IGNORECASE,
)

_WS_RE = re.compile(r"\s+")
_PID_KEYS = ("pid", "process_id", "processId", "ProcessId", "PID")
_TITLE_KEYS = ("title", "window_title", "name", "Title", "WindowTitle")
_HANDLE_KEYS = ("hwnd", "handle", "window_handle", "id", "HWND")


class ReachRefused(ValueError):
    """The reach is refused; the message names the rule and the missing grant."""


# ---- browser history ------------------------------------------------------------


def is_history_database(name: str) -> bool:
    """True for a file name that is a browser history sqlite database."""
    return bool(_HISTORY_DB_RE.match((name or "").strip()))


def history_refusal(rel: str, name: str) -> str:
    return (
        f"{RULE_JAIL}: '{rel}' refused: '{name}' is a browser history database "
        "(Chrome/Edge History, Firefox places.sqlite) and a folder grant does not "
        "admit it; missing grant: history (no such grant kind exists in v1, a "
        "silent History read is out of scope)"
    )


# ---- window attach --------------------------------------------------------------


def normalise_title(title: Any) -> str:
    return _WS_RE.sub(" ", str(title or "")).strip().casefold()


def titles_match(granted: Any, live: Any) -> bool:
    left, right = normalise_title(granted), normalise_title(live)
    return bool(left) and left == right


def _first(row: dict[str, Any], keys: tuple[str, ...]) -> Any:
    for key in keys:
        if key in row and row[key] is not None:
            return row[key]
    return None


def _as_pid(value: Any) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value if value > 0 else None
    try:
        number = int(str(value).strip())
    except (TypeError, ValueError):
        return None
    return number if number > 0 else None


def parse_windows(text: str) -> list[dict[str, Any]]:
    """Normalise a ``list_windows`` result into ``[{pid, title, handle, raw}]``.

    JSON first (a list, or a dict holding a list under ``windows`` / ``result``
    / ``items`` / ``data``). Otherwise one entry per non-empty text line whose
    ``title`` is the whole line and ``pid`` is None: the caller then matches the
    pid as a whole-number token inside the line (``_line_holds``)."""
    body = (text or "").strip()
    if not body:
        return []
    rows: list[dict[str, Any]] = []
    parsed: Any = None
    try:
        parsed = json.loads(body)
    except ValueError:
        parsed = None
    if isinstance(parsed, dict):
        for key in ("windows", "result", "items", "data"):
            if isinstance(parsed.get(key), list):
                parsed = parsed[key]
                break
    if isinstance(parsed, list):
        for item in parsed:
            if not isinstance(item, dict):
                continue
            rows.append(
                {
                    "pid": _as_pid(_first(item, _PID_KEYS)),
                    "title": str(_first(item, _TITLE_KEYS) or ""),
                    "handle": _first(item, _HANDLE_KEYS),
                    "raw": item,
                }
            )
        return rows
    for line in body.splitlines():
        stripped = line.strip()
        if stripped:
            rows.append({"pid": None, "title": stripped, "handle": None, "raw": stripped})
    return rows


def _line_holds(line: str, title: str, pid: int) -> bool:
    if not re.search(rf"(?<!\d){pid}(?!\d)", line):
        return False
    wanted = normalise_title(title)
    return bool(wanted) and wanted in normalise_title(line)


def find_live_window(
    windows: list[dict[str, Any]], title: str, pid: int
) -> dict[str, Any] | None:
    """The live entry whose pid equals ``pid`` and whose title matches ``title``
    under the documented rule, else None."""
    for row in windows:
        live_pid = row.get("pid")
        if live_pid is None:
            if _line_holds(str(row.get("title") or ""), title, pid):
                return row
            continue
        if live_pid == pid and titles_match(title, row.get("title")):
            return row
    return None


def render_attached(grant: dict[str, Any], live: dict[str, Any]) -> str:
    handle = live.get("handle")
    payload = {
        "attached": True,
        "pid": int(grant["pid"]),
        "title": str(grant["title"]),
        "handle": handle,
        "grant_id": grant.get("id"),
        "via": f"{ATTACH_SERVER}.{ATTACH_TOOL}",
        "launched": False,
    }
    return "ATTACHED " + json.dumps(payload, ensure_ascii=True, default=str)


def missing_window_grant(title: str, pid: int | None) -> str:
    return (
        f"{RULE_WINDOW}: no window grant with decision allow for pid={pid} "
        f"title='{title}' in this session; missing grant: window allow (title + pid "
        "as UACC list_windows shows it). Crew never launches OneNote, Word, a PDF "
        "viewer or Explorer: open it yourself and Allow the window"
    )


def window_not_live(title: str, pid: int) -> str:
    return (
        f"{RULE_WINDOW}: window pid={pid} title='{title}' is granted but not open "
        "now (the process exited, the pid changed, or the title changed since the "
        "grant); Crew does not launch or reopen it. Open it yourself and Allow the "
        "window again"
    )


# ---- Excel as data --------------------------------------------------------------


def read_xlsx(
    path: Path,
    rel: str,
    *,
    sheet: str | None = None,
    offset: int = 0,
    limit: int = XLSX_ROWS,
) -> str:
    """Render a workbook's cells as tab-separated rows. Read-only, values only
    (formulas come back as their cached results), never through Excel itself."""
    suffix = path.suffix.lower()
    if suffix not in XLSX_EXTENSIONS:
        allowed = ", ".join(sorted(XLSX_EXTENSIONS))
        raise ReachRefused(f"not a workbook this reader opens as data: '{rel}' ({allowed})")
    if not path.is_file():
        raise ReachRefused(f"not a file: {rel}")
    size = path.stat().st_size
    if size > MAX_XLSX_BYTES:
        raise ReachRefused(
            f"workbook exceeds {MAX_XLSX_BYTES} bytes ({size}); export the sheet you need"
        )
    import zipfile

    import openpyxl
    from openpyxl.utils.exceptions import InvalidFileException

    try:
        book = openpyxl.load_workbook(filename=str(path), read_only=True, data_only=True)
    except (InvalidFileException, zipfile.BadZipFile, OSError, ValueError, KeyError) as exc:
        raise ReachRefused(f"cannot open '{rel}' as a workbook: {exc}") from exc
    try:
        names = list(book.sheetnames)
        if not names:
            raise ReachRefused(f"'{rel}' has no sheets")
        wanted = (sheet or "").strip()
        if wanted and wanted not in names:
            raise ReachRefused(f"sheet '{wanted}' not in '{rel}' (sheets: {', '.join(names)})")
        ws = book[wanted] if wanted else book[names[0]]
        start = max(0, int(offset))
        count = max(1, min(int(limit), XLSX_MAX_ROWS))
        rows: list[str] = []
        total = 0
        for total, values in enumerate(ws.iter_rows(values_only=True), start=1):
            if total <= start:
                continue
            if len(rows) >= count:
                continue
            cells = ["" if v is None else str(v) for v in values[:XLSX_MAX_COLS]]
            rows.append("\t".join(cells).rstrip())
        shown_from = start + 1 if rows else 0
        shown_to = start + len(rows)
        header = (
            f"# {rel} sheet '{ws.title}' rows {shown_from}-{shown_to} of {total} "
            f"(sheets: {', '.join(names)}; read as data via openpyxl, not Excel)\n"
        )
        return header + "\n".join(rows)
    finally:
        book.close()
